from django.shortcuts import render,redirect
from django.contrib.auth import authenticate, login,logout
from django.urls import reverse
from django.http import JsonResponse
from django.http import HttpResponse, HttpResponseRedirect
from movieplanet.models import *
from django.db.models import Q
from django.contrib import messages 
from django.conf import settings
import json
from django.http import JsonResponse
from django.core.serializers.json import DjangoJSONEncoder
import os
import openpyxl
from datetime import datetime
from movieplanet.decorators import (
   permission_required,xhr_request_only
)
######### Admin ##########


def dashboard(request):
    return render(request,"movieplanet/admin/dashboard.html")


@permission_required('Permission')
def permission(request,*args,**kwargs):
   if 'Permission' in kwargs.get('module') and kwargs.get('access'):
      role = kwargs.get('role', '')
      if request.method == 'POST' and request.headers.get('x-requested-with') == 'XMLHttpRequest' and 'View' in kwargs.get('permission'):
            start = request.POST['start']
            length = request.POST['length']
            search = request.POST['search']
            startIndex = (int(start)-1) * int(length)
            endIndex = startIndex + int(length)
            listData = []
            roles = kwargs.get('roleIds')
            if search :
                  data = Role.objects.using('movieplanet').filter(name__icontains=search,id__in=roles)[startIndex:endIndex].all()
                  totalLen = list(Role.objects.using('movieplanet').filter(name__icontains=search,id__in=roles).all())
            else:
                  data = Role.objects.using('movieplanet').filter(id__in=roles)[startIndex:endIndex].all()
                  totalLen = list(Role.objects.using('movieplanet').filter(id__in=roles).all())

            for i in data:
                  permission = {
                  "id":i.id,
                  "roleName":i.name,
                  "action":(f'<a class="btn btn-primary" href="{settings.BASE_URL}movieplanet/admin/administration/permission/{i.id}" >Permission</a>')
                  }
                  listData.append(permission)

            return JsonResponse({
                  "success": True,
                  "iTotalRecords":len(totalLen),
                  "iTotalDisplayRecords":len(totalLen),
                  "aaData":listData
            }, status=200)
      elif role:
          
          if 'Edit' in kwargs.get('permission'):
             module = Module.objects.using('movieplanet').all()
             if request.method == 'POST':
                for m in module:
                    perm = Permission.objects.using('movieplanet').filter(modules_id=m.id,role_id=role).first()
                    if m.parent_id=='':
                        if perm:
                              if m.module in request.POST:
                                perm.permission =  request.POST[m.module]
                                perm.save()
                              else:
                                perm.delete()
                        else:
                              if m.module in request.POST:
                                 Permission.objects.using('movieplanet').create(
                                    permission=request.POST[m.module],
                                    role_id = role,
                                    modules_id = m.id
                                 )
                    else:
                        if m.parent_id:
                              permission = ''
                              if m.module+'[view]' in request.POST:
                                    permission += request.POST[m.module+'[view]']
                              if m.module+'[add]' in request.POST:
                                    permission +=','+request.POST[m.module+'[add]']
                              if m.module+'[edit]' in request.POST:
                                    permission +=','+request.POST[m.module+'[edit]']
                              if m.module+'[delete]' in request.POST:
                                    permission +=','+request.POST[m.module+'[delete]']
                                       
                              if perm:
                                    if permission:
                                       perm.permission =  permission
                                       perm.save()
                                    else:
                                       perm.delete()  
                              else:
                                  if permission:
                                     Permission.objects.using('movieplanet').create(
                                          permission=permission,
                                          role_id = role,
                                          modules_id = m.id,
                                          module_parent_id=m.parent_id
                                     )
                                   
             allow = '<ul class="list-group">'
             for m in module:
                 perm = Permission.objects.using('movieplanet').filter(modules_id=m.id,role_id=role).first()
                 if m.parent_id=='':
                       allow +=(f'<li class="list-group-item ">'
                                f'<div class="w-100 d-flex justify-content-between">'
                                f'<div>{m.module}</div>'
                                f'<div>View <input type="checkbox" name="{m.module}" value="View" {"checked" if perm else None} class="form-check-input" /> </div>'
                                f'</div>'
                                f'{checkParent(role,module,m.id)}'
                                f'</li>')   
                    
             allow +='</ul>'
             return render(request,"movieplanet/admin/allowpermission.html",{'permissions':allow})
          else:
            return redirect(reverse('movieplanet:movieplanet-permission'))
      else:
            return render(request,"movieplanet/admin/permission.html")
   else:
      return render(request,"movieplanet/404.html")  

def checkParent(role,module,mid):
      allow = '<ul class="list-group">'
      for m in module:
            perm = Permission.objects.using('movieplanet').filter(modules_id=m.id,role_id=role).first()
            if m.parent_id == str(mid):
                        checked = []
                        if perm:
                           checked = perm.permission.split(',')
                       
                        allow +=(f'<li class="list-group-item">'
                                    f'<div class="w-100 d-flex justify-content-between">'
                                    f'<div>{m.module}</div>'
                                    f'<div>View <input name="{m.module}[view]" type="checkbox" value="View" {"checked" if "View" in checked else None } class="form-check-input" />'
                                    f'Add <input name="{m.module}[add]" type="checkbox" value="Add" {"checked" if "Add" in checked else None }  class="form-check-input" />' 
                                    f'Edit <input name="{m.module}[edit]" type="checkbox" value="Edit" {"checked" if "Edit" in checked else None } class="form-check-input" />'
                                    f'Delete <input name="{m.module}[delete]" type="checkbox" value="Delete" {"checked" if "Delete" in checked else None } class="form-check-input" /></div>'
                                    f'</div>'
                                    f'{checkParent(role,module,m.id)}'
                                 f'</li>')   
                 
      allow +='</ul>'
      
      return allow




@permission_required('Menu') 
def menu(request,*args,**kwargs):
    if 'Menu' in kwargs.get('module') and kwargs.get('access'):
      if request.method == 'POST' and 'View' in kwargs.get('permission'):
            parentId = kwargs.get('parentId', '')
            start = request.POST['start']
            length = request.POST['length']
            search = request.POST['search']
            startIndex = (int(start)-1) * int(length)
            endIndex = startIndex + int(length)
            file_path = os.path.join(settings.BASE_DIR, 'movieplanet', 'menubar.json')
            listData = []
            with open(file_path, 'r') as file:
                  menus =  json.load(file)
                  m = []
                  if parentId:
                     menus = menuFind(menus,parentId)
                  if search :
                        data = menus[int(startIndex):int(endIndex)]
                  else:
                        data = menus[int(startIndex):int(endIndex)]
                        
                  totalLen = 1
                  
                  for idx in menus:
                        if 'parentId' not in idx:
                              totalLen +=1  
                  for m in data: 
                        name = m['name']
                        id = m['id']
                        if 'Edit' in kwargs.get('permission'):
                              uobj={}
                              uobj['id'] = id
                              uobj['name'] = name
                              uobj['icon'] = m['icon']
                              if m['type']=='file':
                                 uobj['link'] = m['link']
                             
                              btn = f'<button class="btn btn-primary" data-bs-toggle="modal" onclick="openModal({uobj})">Update</button>'
                        else:
                              btn = f'<button class="btn btn-primary">Update</button>' 
                        if parentId:
                              if m['type']=='folder':
                                    name = f'<a href="{settings.BASE_URL}movieplanet/admin/website/menu/{id}">{name}</a>'
                              elif m['type']=='file':
                                    name = f'<a href="#">{name}</a>'
                               
                              post = {
                                    "id":id,
                                    "name":name,
                                    "action":btn
                              }      
                              listData.append(post)
                        else:
                              if 'parentId' not in m:
                                    if m['type']=='folder':
                                          name = f'<a href="{settings.BASE_URL}movieplanet/admin/website/menu/{id}">{name}</a>'
                                    elif m['type']=='file':
                                          name = f'<a href="#">{name}</a>'
                                    post = {
                                          "id":id,
                                          "name":name,
                                          "action":btn
                                    }      
                                    listData.append(post)
            
            action = {}
            if 'Add' in kwargs.get('permission'):
                action['add'] = f'<button class="btn btn-primary" data-bs-toggle="modal" data-bs-target="#AddModel">Add</button>' 
            
            return JsonResponse({
            "success": True,
            "iTotalRecords":totalLen,
            "iTotalDisplayRecords":totalLen,
            "aaData":listData,
            "action":action
            }, status=200)
      elif request.method == 'PUT' and 'Edit' in kwargs.get('permission'):
            file_path = os.path.join(settings.BASE_DIR, 'movieplanet', 'toggle.json')
            data['new_key'] = 'new_value'
            with open(file_path, 'w', encoding='utf-8') as file:
                  json.dump(data, file, indent=4)
      else:
            return render(request,"movieplanet/admin/menu.html")
    else:
      return render(request,"movieplanet/404.html")   

def menuFind(menus, pid):
    result = []
    for m in menus:
        if 'parentId' in m and m['parentId'] == pid:
            result.append(m) 
    return result


@permission_required('Posts') 
def posts(request,*args,**kwargs):
    if 'Posts' in kwargs.get('module') and kwargs.get('access'):
      parentId = kwargs.get('parentId', None)
      if request.method == 'POST' and 'View' in kwargs.get('permission'):
            start = request.POST['start']
            length = request.POST['length']
            search = request.POST['search']
            startIndex = (int(start)-1) * int(length)
            endIndex = startIndex + int(length)
            
            if search :
                  data = Posts.objects.using('movieplanet').filter(Q(parent=parentId),name__icontains=search,status=1)[startIndex:endIndex].all()
                  totalLen = Posts.objects.using('movieplanet').filter(Q(parent=parentId),name__icontains=search,status=1).count()
            
            else:
                  data = Posts.objects.filter(Q(parent=parentId),status=1).all()[startIndex:endIndex]
                  totalLen = Posts.objects.filter(Q(parent=parentId),status=1).count()
            
            listData = []
            for i in data:
                  btn =''
                  if 'Edit' in kwargs.get('permission'):
                     obj = {
                          "id":i.id,
                          "name":i.name,
                          "image":i.image,
                          "rate":i.rate,
                          "menu":i.menu if i.menu else '',
                          "size":i.size,
                          "lang":i.lang,
                          "genre":i.genre,
                          "story":i.story if i.story else '' ,
                          "link":i.link if i.link else '' ,
                          "starcast":i.starcast if i.starcast else '',
                          "release_date":i.release_date if i.release_date else '',
                          "type":i.type,
                          "status":i.status
                     }
                     btn += f'<button class="btn btn-primary" onclick="openModal({obj})" >Edit</button>'
                  if 'Delete' in kwargs.get('permission'):
                     btn += f'<button class="btn btn-primary">Delete</button>'
                  
                  link = i.name
                  if i.type==2:
                     link = f'<a href="{settings.BASE_URL}movieplanet/admin/website/posts/{i.id}" >{i.name}</a>'
                  post = {
                        "id":i.id,
                        "name":link,
                        "image":f'<img src={i.image}/>',
                        "rate":i.rate,
                        "action":btn
                  }
                        
                  listData.append(post)
            action = {}
            if 'Add' in kwargs.get('permission'):
               action['add'] = f'<button class="btn btn-primary" onclick="openModal()">Add</button>'

            return JsonResponse({
                  "success": True,
                  "iTotalRecords":totalLen,
                  "iTotalDisplayRecords":totalLen,
                  "aaData":listData,
                  "action":action
            }, status=200)
      elif request.method == 'PUT' and 'Add' in kwargs.get('permission'):
            try:
                  body_unicode = request.body.decode('utf-8')
                  if not body_unicode:
                        return JsonResponse({"error": "Empty request body"}, status=400)
                  post = json.loads(body_unicode)
                  if Posts.objects.filter(name=post['name']).exists():
                       msg="Movie exist"
                  else:
                        Posts.objects.create(
                              name=post['name'],
                              image=post.get('image', ''),
                              rate=post.get('rate', 'N/A'),
                              size=post.get('size', 'N/A'),
                              genre=post.get('genre', 'N/A'),
                              type=post.get('type', 2),
                              lang=post.get('lang', 'N/A'),
                              story=post.get('story', 'N/A'),
                              status=post.get('status', 0),
                              link=post.get('link', ''),
                              parent=parentId
                        )
                        msg="Inserted success"
                  return JsonResponse({
                        "status":True,
                        "message":msg
                  })
            
            except json.JSONDecodeError:
                  return JsonResponse({"error": "Invalid JSON format"}, status=400)
      else:
            return render(request,"movieplanet/admin/post.html")
    else:
      return render(request,"movieplanet/404.html")  


def excelPost(request,*args,**kwargs):
      try:
            if request.FILES["excel_file"]:
                  excel_file = request.FILES["excel_file"]
                  wb = openpyxl.load_workbook(excel_file)
                  # sheets = wb.sheetnames
                  worksheet = wb["Posts"]
                  # active_sheet = wb.active
                  excel_data = list()
                  for i, row in enumerate(worksheet.iter_rows()):
                        if i == 0:
                           continue 
                        row_data = list()
                        for cell in row:
                              row_data.append(str(cell.value))
                        excel_data.append(row_data)
                  posts = []
                  for row in excel_data:
                        posts.append(Posts(
                              name=row[0],
                              rate=int(row[1]),
                              size=int(row[2]),
                              lang=row[3],
                              image=row[4],
                              genre=row[5],
                              link=row[6],
                              starcast=row[7],
                              type=int(row[8]),
                              status=int(row[9]),
                              # release_date=datetime.strptime(raw_date, "%d-%m-%Y %H:%M:%S"),
                              story=row[11]
                        ))           
                  Posts.objects.bulk_create(posts)
                  return JsonResponse({
                  "status":True,
                  "message":"Inserted success"
                  })
      except json.JSONDecodeError:
            return JsonResponse({"error": "Invalid JSON format"}, status=400)   

@permission_required('Users')    
def customers(request,*args,**kwargs):
   if 'Users' in kwargs.get('module') and kwargs.get('access'):
      if request.method == 'POST' and 'View' in kwargs.get('permission'):
            start = request.POST['start']
            length = request.POST['length']
            search = request.POST['search']
            startIndex = (int(start)-1) * int(length)
            endIndex = startIndex + int(length)

            # print(kwargs.get('roleIds'))
            listData = []
            totalLen=0
            
            excIds = list(Roles.objects.using('movieplanet').filter(user_id=kwargs.get('authId')).values_list('given_id',flat=True))
            excIds.append(kwargs.get('authId'))
            if search :
                  data = Customer.objects.using('movieplanet').filter(is_admin="0",name__icontains=search).exclude(id__in=excIds)[startIndex:endIndex].all()
                  totalLen = Customer.objects.using('movieplanet').filter(is_admin="0",name__icontains=search).exclude(id__in=excIds).count()
            else:
                  data = Customer.objects.using('movieplanet').filter(is_admin="0").exclude(id__in=excIds)[startIndex:endIndex].prefetch_related('roles')
                  totalLen = Customer.objects.using('movieplanet').filter(is_admin="0").exclude(id__in=excIds).count()
            
            
            for i in data:
                  
                  if 'Edit' in kwargs.get('permission'):
                        rol=[]
                        for r in i.roles.all():
                              urol={
                                    "user_id":r.user_id,
                                    "role":r.role.name,
                                    "assign":r.assign if r.assign else ''
                              }
                              rol.append(urol)
                        editbtn = f'<button class="btn btn-sm  btn-danger" onclick="openModal({rol},{i.id})" >Role</button>'
                  else:
                        editbtn = ''
                  roles = list(i.roles.values_list('role_id',flat=True))
                  if kwargs.get('roleIds') == roles or roles==[] or kwargs.get('isAdmin') or all(item in kwargs.get('roleIds') for item in roles):
                        obj = {
                              "id":i.id,
                              "name":i.name,
                              "email":i.email,
                              "action":editbtn
                        }  
                        listData.append(obj)

            role = Role.objects.using('movieplanet').filter(id__in=kwargs.get('roleIds')).prefetch_related('role')
            
            action = []
            
            for r in role:
                if kwargs.get('isAdmin'): 
                  obj = {
                        "id":r.id,
                        "role":r.name,
                        "assign":True
                  }
                  action.append(obj)
                elif r.role.filter(role_id__in=kwargs.get('roleIds'),user_id=kwargs.get('authId'),assign="1").exists():
                     obj = {
                        "id":r.id,
                        "role":r.name,
                        "assign":True
                     }
                     action.append(obj)
                  

            return JsonResponse({
                        "success": True,
                        "iTotalRecords":totalLen,
                        "iTotalDisplayRecords":totalLen,
                        "aaData":listData,
                        "action":action
                  }, status=200)
      elif  request.method == 'PUT' and 'Edit' in kwargs.get('permission'):
            try:
                  body_unicode = request.body.decode('utf-8')
                  if not body_unicode:
                        return JsonResponse({"error": "Empty request body"}, status=400)
                  post = json.loads(body_unicode)
                  
                  role = Role.objects.using('movieplanet').filter(id__in=kwargs.get('roleIds')).prefetch_related('role')
                  
                  action = []
                  for r in role:
                        if r.name in post:
                           
                              if kwargs.get('isAdmin'): 
                                    if r.role.filter(role_id=post[r.name],user_id=post['customer']).exists():
                                       robj = r.role.filter(role_id=post[r.name],user_id=post['customer']).first()
                                       if "assign"+r.name in post:
                                          robj.assign = "1"
                                       else:
                                          robj.assign = ""
                                       robj.given_id = kwargs.get('authId')
                                       robj.save()        
                                    else:
                                       Roles.objects.create(
                                          role_id=post[r.name],
                                          user_id = post['customer'],
                                          assign =  "1" if post.get("assign" + r.name) else "",
                                          given_id = kwargs.get('authId')
                                       )
                              elif r.role.filter(role_id=post[r.name],user_id=kwargs.get('authId'),assign="1").exists():
                                    if r.role.filter(role_id=post[r.name],user_id=post['customer']).exists():
                                       robj = r.role.filter(role_id=post[r.name],user_id=post['customer']).first()
                                       if "assign"+r.name in post:
                                          robj.assign = "1"
                                       else:
                                          robj.assign = ""
                                       robj.given_id = kwargs.get('authId')
                                       robj.save()        
                                    else:
                                       pass                                       
                        else:
                              Roles.objects.filter(
                              role_id=r.id,
                              user_id=post.get('customer'),
                              ).delete()        
                                         

                  return JsonResponse({
                  "status":True,
                  "message":"Inserted success"
                  })
            
            except json.JSONDecodeError:
                  return JsonResponse({"error": "Invalid JSON format"}, status=400)
      return render(request,"movieplanet/admin/user.html")
      # return HttpResponseRedirect(request.META['HTTP_REFERER']) 
   else:
      return render(request,"movieplanet/404.html")  


    
 
@xhr_request_only()
def sidebarList(request,*args,**kwargs):
   modules = kwargs.get('module')
   sidebarList =  list(Module.objects.using('movieplanet').filter(module__in=modules).values())
   return JsonResponse({
      "success": True,
      "data":sidebarList
   }, status=200)   



############# Auth ############
def login(request):
    if request.method == 'POST':
        email = request.POST['email']
        password = request.POST['password']
        customer = Customer.authenticate(email=email, password=password)
        if customer:
            request.session['customer'] = json.loads(
                  json.dumps({
                        'id': customer.id,
                        'name': customer.name,
                        'email': customer.email,
                        'is_admin':customer.is_admin
                  }, cls=DjangoJSONEncoder)
            )
            return redirect('admin/dashboard')
            # return HttpResponse(f"Welcome {customer.name}!")
        else:
            return HttpResponseRedirect(reverse('movieplanet:movieplanet-login'))
    else:
      return render(request,"movieplanet/login.html")
   
def signup(request):
    if request.method == 'POST':
        email = request.POST['email']
        name = request.POST['name']
        password = request.POST['password']
        if Customer.objects.filter(email=email).exists():
            messages.error(request, "Email already exists. Please log in or use another email.")
            return HttpResponseRedirect(reverse('movieplanet-signup'))
        customer = Customer(email=email, name=name,is_admin=0)
        customer.set_password(password)
        customer.save()
        return redirect('admin/dashboard')
        # return HttpResponseRedirect(reverse('admin/dashboard')) 
    else:
      return render(request,"movieplanet/signup.html")

def logout(request):
    request.session.flush()
    return HttpResponseRedirect(reverse('movieplanet:movieplanet-login'))
    # return HttpResponse("Logged out successfully!")



########## Frontend ################

def home(request,*args,**kwargs):
    if request.method == 'POST':
      start = request.POST['start']
      length = request.POST['length']
      search = request.POST['search']
      startIndex = (int(start)-1) * int(length)
      endIndex = startIndex + int(length)
      Link = kwargs.get('Link', None)
      if Link:
            linkList = Link.split("+")
            Link = " ".join(linkList)
            parent = Posts.objects.filter(name=Link,status=1).first()
            Link = parent.id
     
      if search :
            data = Posts.objects.filter(Q(parent=Link),name__icontains=search,status=1)[startIndex:endIndex].all()
            totalLen = Posts.objects.filter(Q(parent=Link),name__icontains=search,status=1).count()
      
      else:
            data = Posts.objects.filter(Q(parent=Link),status=1)[startIndex:endIndex].all()
            totalLen = Posts.objects.filter(Q(parent=Link),status=1).count()
      
      listData = []
      for i in data:
            post = {
                  "id":i.id,
                  "name":i.name,
                  "image":i.image,
                  "rate":i.rate,
                  "type":i.type
            }     
            listData.append(post)
      
      return JsonResponse({
      "success": True,
      "iTotalRecords":totalLen,
      "iTotalDisplayRecords":totalLen,
      "aaData":listData
      }, status=200)
    else:
      return render(request,"movieplanet/home.html")



def category(request,*args,**kwargs):
    
    if request.method == 'POST':
      start = request.POST['start']
      length = request.POST['length']
      search = request.POST['search']
      startIndex = (int(start)-1) * int(length)
      endIndex = startIndex + int(length)
      params  = kwargs.get('params')
      categories = params.replace("/", " ").split()
      query = Q()

      for word in categories:
          query &= Q(menu__icontains=word)
      if search :
            data = Posts.objects.filter(query,parent=None,name__icontains=search,status=1)[startIndex:endIndex].all()
            totalLen = Posts.objects.filter(query,parent=None,name__icontains=search,status=1).count()
      
      else:
            data = Posts.objects.filter(query,parent=None,status=1)[startIndex:endIndex].all()
            totalLen = Posts.objects.filter(query,parent=None,status=1).count()
      
      listData = []
      for i in data:
            post = {
                  "id":i.id,
                  "name":i.name,
                  "image":i.image,
                  "rate":i.rate,
                  "type":i.type
            }     
            listData.append(post)
      
      return JsonResponse({
      "success": True,
      "iTotalRecords":totalLen,
      "iTotalDisplayRecords":totalLen,
      "aaData":listData
      }, status=200)
    else:
      return render(request,"movieplanet/home.html")



def detail(request,Link=None):
      linkList = Link.split("+")
      MovieName = " ".join(linkList)
      data = Posts.objects.filter(name=MovieName,status=1).values().first()
      context = {
         "link":Link,
         "post":data
      }
      return render(request,"movieplanet/detail.html",context)



def menubar(request,*args,**kwargs):
      file_path = os.path.join(settings.BASE_DIR, 'movieplanet', 'menubar.json')
      try:
            with open(file_path, 'r') as file:
                  menus =  json.load(file)
                  MenuHtml = ''
                  for m in menus:
                        
                        if 'parentId' not in m:
                              if m['type'] == 'folder':
                                    MenuHtml += (f'<li><button class="nav-link dropdown-btn" data-dropdown="dropdown{m["id"]}" aria-haspopup="true" aria-expanded="false" aria-label="discover">{m["name"]}<i class="bx bx-chevron-down" aria-hidden="true"></i></button>'
                                                f'{menuLoop(menus,m["id"])}'
                                                f'</li>')
                              elif m['type']=='file':  
                                    MenuHtml += (f'<li><a class="nav-link dropdown-link dropdown-btn" data-dropdown="dropdown{m["id"]}" href="{m["link"]}" aria-haspopup="true" aria-expanded="false">{m["name"]}</a></li>')

                              
                  
            return JsonResponse({"status":True,"Menus":MenuHtml})
      except Exception as e:
            return JsonResponse({"status":False,"error": str(e)}, status=500)


def menuLoop(Menus=[],MenuId=None,IsLoop=None):
   check = False
   menu = f'<div id="dropdown{MenuId}" class="dropdown"><ul role="menu">'
   arrow = ''
   if IsLoop:
      menu = f'<div id="dropdown{MenuId}" class="dropdown loopMenu"><ul role="menu">'
      arrow = '<i class="bx bx-chevron-down" aria-hidden="true"></i>'
   for m in Menus:
      check = True
      
      if 'parentId' in m and m['parentId'] == MenuId:
            if m['type'] == 'folder':
                  menu += (f'<li><button class="nav-link dropdown-btn" data-dropdown="dropdown{m["id"]}" aria-haspopup="true" aria-expanded="false" aria-label="discover">{m["name"]}<i class="bx bx-chevron-down" aria-hidden="true"></i></button>'
                              f'{menuLoop(Menus,m["id"],True)}'
                              f'</li>')
            elif m['type']=='file':  
                 menu += (f'<li><a class="nav-link dropdown-link dropdown-btn" data-dropdown="dropdown{m["id"]}" href="{m["link"]}" aria-haspopup="true" aria-expanded="false">{m["name"]}</a></li>')

                                   
   menu +='</ul></div>'
   if check:
      return menu
   else:    
      return ''





"""
def menu(request,*args,**kwargs):
    parentId = kwargs.get('parentId', '')
    if request.method == 'POST':
      start = request.POST['start']
      length = request.POST['length']
      search = request.POST['search']
      startIndex = (int(start)-1) * int(length)
      endIndex = startIndex + int(length)
      file_path = os.path.join(settings.BASE_DIR, 'movieplanet', 'sidebar.json')
      listData = []
      with open(file_path, 'r') as file:
            menus =  json.load(file)
            m = []
            if parentId:
               menus = menuFind(menus,parentId)
            print(m)
            if search :
                  data = menus[int(startIndex):int(endIndex)]
                  totalLen = len(menus[int(startIndex):int(endIndex)])
            else:
                  
                  data = menus[int(startIndex):int(endIndex)]
                  totalLen = len(menus[int(startIndex):int(endIndex)])
      
            # print(d)
            for i in data:
                  name = i['name']
                  id = i['id']
                  if i['type']=='folder':
                     name = f'<a href="{settings.BASE_URL}movieplanet/admin/website/menu/{id}">{name}</a>'

                  post = {
                        "id":id,
                        "name":name,
                        "action":'<button class="btn btn-primary">Update</button>'
                  }      
                  listData.append(post)
      
      return JsonResponse({
      "success": True,
      "iTotalRecords":totalLen,
      "iTotalDisplayRecords":totalLen,
      "aaData":listData
      }, status=200)
    else:
      return render(request,"movieplanet/admin/menu.html")

def menuFind(menus, idx):
    result = []

    def recurse(items, id):
        nonlocal result  # to modify outer variable
        for item in items:
            if item.get('id') == id and item.get('type') == 'folder':
                result = item.get('children', [])
                return True  # found, stop searching
            if 'children' in item and isinstance(item['children'], list):
                if recurse(item['children'], id):
                    return True  # bubble up success
        return False  # not found in this branch

    recurse(menus, idx)
    return result

def sidebarList(request,*args,**kwargs):
      file_path = os.path.join(settings.BASE_DIR, 'movieplanet', 'sidebar.json')
      try:
            with open(file_path, 'r') as file:
                  menus =  json.load(file)
                  MenuHtml = ''
                  for m in menus:
                        
                        if m['type'] == 'folder':
                              MenuHtml += (f'<li><button class="nav-link dropdown-btn" data-dropdown="dropdown{m["id"]}" aria-haspopup="true" aria-expanded="false" aria-label="discover">{m["name"]}<i class="bx bx-chevron-down" aria-hidden="true"></i></button>'
                                          f'{menuLoop(m["children"],m["id"])}'
                                          f'</li>')
                        elif m['type'] == 'file':
                              MenuHtml += (f'<li><a class="nav-link dropdown-link dropdown-btn" data-dropdown="dropdown{m["id"]}" href="#" aria-haspopup="true" aria-expanded="false">{m["name"]}</a></li>')
                  
            return JsonResponse({"status":True,"Menus":MenuHtml})
      except Exception as e:
            return JsonResponse({"status":False,"error": str(e)}, status=500)


def menuLoop(Menus=[],MenuId=None,IsLoop=None):
   check = False
   if IsLoop:
      menu = f'<div id="dropdown{MenuId}" class="dropdown loopMenu"><ul role="menu">'
      arrow = '<i class="bx bx-chevron-down" aria-hidden="true"></i>'
   else:
      menu = f'<div id="dropdown{MenuId}" class="dropdown"><ul role="menu">'
      arrow = ''
   for m in Menus:
      check = True
      if m['type']=='folder':
            menu += (f'<li><a class="nav-link dropdown-link dropdown-btn" data-dropdown="dropdown{m["id"]}" href="#" aria-haspopup="true" aria-expanded="false">{m["name"]}{arrow}</a>{menuLoop(m["children"],m["id"],True)}</li>')
      elif m['type'] == 'file':
            
            menu += (f'<li><a class="nav-link dropdown-link dropdown-btn" data-dropdown="dropdown{m["id"]}" href="#" aria-haspopup="true" aria-expanded="false">{m["name"]}</a></li>')
               
   menu +='</ul></div>'
   if check:
      return menu
   else:    
      return ''
   

def menuLoop(Menus=[],MenuId=None,IsLoop=None):
   check = False
   if IsLoop:
      menu = f'<div id="dropdown{MenuId}" class="dropdown loopMenu"><ul role="menu">'
      arrow = '<i class="bx bx-chevron-down" aria-hidden="true"></i>'
   else:
      menu = f'<div id="dropdown{MenuId}" class="dropdown"><ul role="menu">'
      arrow = ''
   for m in Menus:
      if m.menuId:
         if m.menuId==str(MenuId):
            check = True
            menu += (f'<li><a class="dropdown-link dropdown-btn" data-dropdown="dropdown{m.id}" href="#" aria-haspopup="true" aria-expanded="false">{m.name}{arrow}</a>{menuLoop(Menus,m.id,True)}</li>')
   menu +='</ul></div>'
   if check:
      return menu
   else:    
      return ''

def menuList(request):
   if request.method == 'POST':
      menus =  Menu.objects.filter(status="1").all()
      MenuHtml = ''
      for m in menus:
         if m.menuId == None:
            MenuHtml += (f'<li><button class="nav-link dropdown-btn" data-dropdown="dropdown{m.id}" aria-haspopup="true" aria-expanded="false" aria-label="discover">{m.name}<i class="bx bx-chevron-down" aria-hidden="true"></i></button>'
                         f'{menuLoop(menus,m.id)}'
                         f'</li>')
      return JsonResponse({
               "Success": True,
               "Menus":MenuHtml
            }, status=200)

def menu(request):

    if request.method == 'POST':
      start = request.POST['start']
      length = request.POST['length']
      search = request.POST['search']
      startIndex = (int(start)-1) * int(length)
      endIndex = startIndex + int(length)
      
      
      if search :
            data = Menu.objects.filter(name__icontains=search,status=1)[startIndex:endIndex].all()
            totalLen = Menu.objects.filter(name__icontains=search,status=1).count()
      
      else:
            data = Menu.objects.filter(status=1).all()[startIndex:endIndex]
            totalLen = Menu.objects.filter(status=1).count()
      
      listData = []
      for i in data:
            post = {
                  "id":i.id,
                  "name":i.name,
                  "status":i.status
            }
                  
            listData.append(post)
      
      return JsonResponse({
      "success": True,
      "iTotalRecords":totalLen,
      "iTotalDisplayRecords":totalLen,
      "aaData":listData
      }, status=200)
    else:
        return render(request,"movieplanet/admin/menu.html") 
"""